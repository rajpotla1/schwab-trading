import importlib.util
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock, patch

import openpyxl
import requests


spec = importlib.util.spec_from_file_location(
    "positions_sync", Path(__file__).resolve().parents[1] / "sync_orders_to_dashboard.py"
)
sync = importlib.util.module_from_spec(spec)
spec.loader.exec_module(sync)


class PositionSyncTests(unittest.TestCase):
    def test_workbook_opens_only_after_successful_save(self):
        for save_error, open_error in (
            (None, None), (PermissionError("locked"), None),
            (None, OSError("no associated app")),
        ):
            with self.subTest(save_error=save_error, open_error=open_error):
                events = []
                def save(*args, **kwargs):
                    events.append("save")
                    if save_error:
                        raise save_error

                def launch(*args):
                    events.append("open")
                    if open_error:
                        raise open_error

                with patch.object(sync, "load_dotenv"), patch.dict(sync.os.environ, {
                    "SCHWAB_CLIENT_ID": "test", "SCHWAB_CLIENT_SECRET": "test",
                    "SCHWAB_CALLBACK_URL": "https://localhost",
                }), patch.object(sync.schwabdev, "Client"), \
                     patch.object(Path, "is_file", return_value=True), \
                     patch.object(sync, "fetch_positions", return_value=[]), \
                     patch.object(sync, "fetch_open_orders", return_value=[]), \
                     patch.object(sync, "write_positions", side_effect=save), \
                     patch.object(sync.os, "startfile", side_effect=launch) as opener:
                    if save_error or open_error:
                        with self.assertRaises(SystemExit) as error:
                            sync.main()
                        if open_error:
                            self.assertIn("Workbook saved, but", str(error.exception))
                    else:
                        sync.main()
                    if save_error:
                        opener.assert_not_called()
                        self.assertEqual(events, ["save"])
                    else:
                        opener.assert_called_once_with(sync.MASTER_FILENAME)
                        self.assertEqual(events, ["save", "open"])

    def test_open_orders_include_both_sides_and_exclude_terminal(self):
        client = Mock()
        client.linked_accounts.return_value.json.return_value = [
            {"hashValue": "hash", "accountNumber": "123456789"}
        ]
        def order(order_id, status="WORKING", instruction="SELL"):
            return {
                "orderId": order_id, "status": status, "orderType": "LIMIT",
                "price": 300, "remainingQuantity": 1, "filledQuantity": 1,
                "orderLegCollection": [{
                    "quantity": 2, "instruction": instruction,
                    "instrument": {"symbol": "IQV"},
                }],
            }
        parent = order(1, "FILLED", "BUY")
        parent["childOrderStrategies"] = [
            order(2, "AWAITING_PARENT_ORDER"),
            order(6, "AWAITING_PARENT_ORDER", "BUY_TO_COVER"),
        ]
        client.account_orders.return_value.json.return_value = [
            parent, order(3, "QUEUED"), order(4, "CANCELED"), order(5, instruction="BUY")
        ]
        rows = sync.fetch_open_orders(client)
        self.assertEqual([r[12] for r in rows], ["2", "3", "5", "6"])
        self.assertEqual([r[12] for r in rows if r[2].startswith("SELL")], ["2", "3"])
        self.assertEqual([r[12] for r in rows if r[2].startswith("BUY")], ["5", "6"])
        self.assertEqual(rows[0][3:6], (2, 1, 1))
        calls = client.account_orders.call_args_list
        self.assertEqual(len(calls), 3)
        from datetime import datetime
        start = datetime.fromisoformat(calls[-1].kwargs["fromEnteredTime"])
        end = datetime.fromisoformat(calls[0].kwargs["toEnteredTime"])
        self.assertEqual((end - start).total_seconds(), 90 * 24 * 60 * 60)
        for call in calls:
            self.assertEqual(call.kwargs["maxResults"], 3000)
            self.assertNotIn("status", call.kwargs)
        client.place_order.assert_not_called()
        client.replace_order.assert_not_called()

    def test_failed_or_truncated_sell_query_aborts(self):
        for truncated in (False, True):
            client = Mock()
            client.linked_accounts.return_value.json.return_value = [{"hashValue": "hash"}]
            if truncated:
                client.account_orders.return_value.json.return_value = [{}] * 3000
                error = ValueError
            else:
                client.account_orders.return_value.raise_for_status.side_effect = requests.HTTPError("401")
                error = requests.HTTPError
            with self.subTest(truncated=truncated), self.assertRaises(error):
                sync.fetch_open_orders(client)

    def test_sell_sheet_refresh_clears_stale_orders_and_keeps_other_sheets(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.xlsx"
            workbook = openpyxl.Workbook()
            workbook.active.title = "Other"
            workbook["Other"]["A1"] = "keep"
            workbook.create_sheet(sync.SELL_SHEET)["A5"] = "old order"
            workbook.save(path)
            workbook.close()
            row = ("account", "IQV", "SELL", 2, 0, 2, "LIMIT", 300, None,
                   "WORKING", "GOOD_TILL_CANCEL", "NORMAL", "123", "time", "SINGLE")
            buy_row = (*row[:2], "BUY", *row[3:12], "456", *row[13:])
            sync.write_positions([], path, sell_rows=[row], buy_rows=[buy_row])
            workbook = openpyxl.load_workbook(path)
            try:
                sheet = workbook[sync.SELL_SHEET]
                self.assertEqual(tuple(c.value for c in sheet[1]), sync.SELL_COLUMNS)
                self.assertEqual(sheet["M2"].value, "123")
                self.assertIsNone(sheet["A5"].value)
                self.assertEqual(workbook["Other"]["A1"].value, "keep")
                buy_sheet = workbook[sync.BUY_SHEET]
                self.assertEqual(tuple(c.value for c in buy_sheet[1]), sync.SELL_COLUMNS)
                self.assertEqual(buy_sheet["C2"].value, "BUY")
                self.assertEqual(buy_sheet["M2"].value, "456")
                self.assertEqual(buy_sheet.auto_filter.ref, "A1:O2")
            finally:
                workbook.close()
            sync.write_positions([], path, sell_rows=[], buy_rows=[])
            workbook = openpyxl.load_workbook(path)
            try:
                self.assertIsNone(workbook[sync.SELL_SHEET]["A2"].value)
                self.assertIsNone(workbook[sync.BUY_SHEET]["A2"].value)
            finally:
                workbook.close()

    def test_exports_require_macro_free_format(self):
        self.assertEqual(sync.MASTER_FILENAME.suffix, ".xlsx")
        with self.assertRaisesRegex(ValueError, "macro-free"):
            sync.write_positions([], "dashboard.xlsm")

    def client(self, data):
        client = Mock()
        client.account_details_all.return_value.json.return_value = data
        client.quotes.return_value.json.return_value = {
            "TMO": {
                "quote": {"lastPrice": 999},
                "regular": {"regularMarketLastPrice": 600,
                            "regularMarketPercentChange": 2, "regularMarketNetChange": 12},
            }
        }
        return client

    def test_position_fields_and_quote_metrics_are_exported(self):
        client = self.client([{"securitiesAccount": {
            "accountNumber": "123456789",
            "positions": [{
                "instrument": {"symbol": "TMO", "description": "Company", "assetType": "EQUITY"},
                "longQuantity": 2, "shortQuantity": 0,
                "averagePrice": 450, "marketValue": 1200, "longOpenProfitLoss": 200,
            }],
        }}])
        client.quotes.return_value.json.return_value["TMO"]["quote"]["lastPrice"] = 999
        self.assertEqual(sync.fetch_positions(client), [
            ("TMO - Company", 2, 600, 0.02, 12, 1200, 0.2, 200, 500, 1000)
        ])
        client.account_details_all.assert_called_once_with(fields="positions")

    def test_missing_reported_profit_is_not_recomputed_from_average_price(self):
        client = self.client([{"accountNumber": "123", "positions": [{
            "instrument": {"symbol": "TMO", "assetType": "EQUITY"},
            "longQuantity": 2, "shortQuantity": 0,
            "averagePrice": 500, "marketValue": 1200,
        }]}])
        row = sync.fetch_positions(client)[0]
        self.assertEqual(row[2], 600)
        self.assertEqual(row[6:], (None, None, None, None))

    def test_http_error_is_not_treated_as_empty_positions(self):
        client = self.client({"error": "unauthorized"})
        client.account_details_all.return_value.raise_for_status.side_effect = requests.HTTPError("401")
        with self.assertRaises(requests.HTTPError):
            sync.fetch_positions(client)

    def test_short_and_zero_cost_positions(self):
        client = self.client([{"accountNumber": "123", "positions": [{
            "instrument": {"symbol": "TMO", "assetType": "EQUITY"},
            "longQuantity": 0, "shortQuantity": 2,
            "averagePrice": 700, "marketValue": -1200, "shortOpenProfitLoss": 200,
        }]}])
        row = sync.fetch_positions(client)[0]
        self.assertEqual(row[1], -2)
        self.assertEqual(row[7:], (200, 700, -1400))
        self.assertAlmostEqual(row[6], 200 / 1400)
        client.account_details_all.return_value.json.return_value[0]["positions"][0]["averagePrice"] = 0
        client.account_details_all.return_value.json.return_value[0]["positions"][0]["shortOpenProfitLoss"] = -1200
        row = sync.fetch_positions(client)[0]
        self.assertIsNone(row[6])
        self.assertEqual(row[9], 0)

    def test_missing_quotes_are_blank_not_fabricated(self):
        client = self.client([{"accountNumber": "123", "positions": [{
            "instrument": {"symbol": "TMO", "assetType": "EQUITY"},
            "longQuantity": 2, "shortQuantity": 0,
            "averagePrice": 500, "marketValue": 1200,
        }]}])
        client.quotes.return_value.json.return_value = {}
        row = sync.fetch_positions(client)[0]
        self.assertEqual(row[2:5], (None, None, None))
        self.assertIsNone(row[5])

    def test_quote_http_failure_aborts_export(self):
        client = self.client([{"accountNumber": "123", "positions": [{
            "instrument": {"symbol": "TMO", "assetType": "EQUITY"}
        }]}])
        client.quotes.return_value.raise_for_status.side_effect = requests.HTTPError("503")
        with self.assertRaises(requests.HTTPError):
            sync.fetch_positions(client)

    def test_positions_sorted_by_gain_percent_descending_with_blanks_last(self):
        positions = [{
            "instrument": {"symbol": symbol, "assetType": "EQUITY"},
            "longQuantity": 1, "shortQuantity": 0,
            "averagePrice": cost, "marketValue": value,
            "longOpenProfitLoss": value - cost,
        } for symbol, cost, value in (
            ("LOSS", 100, 90), ("UNAVAILABLE", 0, 10),
            ("SMALL", 100, 105), ("LARGE", 100, 120),
        )]
        client = self.client([{"accountNumber": "123", "positions": positions}])
        client.quotes.return_value.json.return_value = {
            p["instrument"]["symbol"]: {
                "regular": {
                    "regularMarketLastPrice": p["marketValue"],
                    "regularMarketPercentChange": 0, "regularMarketNetChange": 0,
                }
            } for p in positions
        }
        rows = sync.fetch_positions(client)
        self.assertEqual([r[0] for r in rows], ["LARGE", "SMALL", "LOSS", "UNAVAILABLE"])

    def test_iqv_matches_schwab_web_screenshot(self):
        client = self.client([{"accountNumber": "123", "positions": [{
            "instrument": {"symbol": "IQV", "assetType": "EQUITY"},
            "longQuantity": 1, "shortQuantity": 0, "averagePrice": 266.73,
            "marketValue": 272.25, "longOpenProfitLoss": 5.52,
        }]}])
        client.quotes.return_value.json.return_value = {"IQV": {
            "quote": {"lastPrice": 265.3157, "netChange": -1.5843},
            "extended": {"mark": 272.25},
            "reference": {"description": "IQVIA HLDGS INC"},
            "regular": {"regularMarketLastPrice": 266.9,
                        "regularMarketNetChange": -4.14,
                        "regularMarketPercentChange": -1.52744982},
        }}
        row = sync.fetch_positions(client)[0]
        self.assertEqual(row[0], "IQV - IQVIA HLDGS INC")
        self.assertEqual(row[2], 266.9)
        self.assertEqual(row[5], 266.9)
        self.assertEqual(row[4], -4.14)
        self.assertAlmostEqual(row[3], -0.0152744982)
        self.assertAlmostEqual(row[7], 0.17)
        self.assertEqual(f"{row[6]:.2%}", "0.06%")
        self.assertAlmostEqual(row[9], 266.73)

    def test_mutual_fund_uses_nav(self):
        client = self.client([{"accountNumber": "123", "positions": [{
            "instrument": {"symbol": "SWVXX", "assetType": "MUTUAL_FUND"},
            "longQuantity": 100, "shortQuantity": 0,
            "marketValue": 100, "longOpenProfitLoss": 0,
        }]}])
        client.quotes.return_value.json.return_value = {"SWVXX": {
            "quote": {"nAV": 1, "netChange": 0, "netPercentChange": 0}
        }}
        row = sync.fetch_positions(client)[0]
        self.assertEqual(row[2:8], (1, 0, 0, 100, 0, 0))

    def test_invalid_or_missing_positions_abort(self):
        for data in ([], {"error": "failure"}, [{"accountNumber": "123"}],
                     [{"accountNumber": "123", "positions": None}],
                     [{"accountNumber": "123", "positions": [{}]}]):
            with self.subTest(data=data), self.assertRaises(ValueError):
                sync.fetch_positions(self.client(data))

    def test_explicit_empty_positions_are_valid(self):
        self.assertEqual(sync.fetch_positions(self.client([
            {"accountNumber": "123", "positions": []}
        ])), [])

    def test_only_target_sheets_are_refreshed_and_stale_rows_cleared(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.xlsx"
            workbook = openpyxl.Workbook()
            workbook.active.title = "Other"
            workbook["Other"]["A1"] = "=1+1"
            for name in sync.SHEET_NAMES:
                sheet = workbook.create_sheet(name)
                sheet["A10"] = "obsolete"
                sheet["N2"] = "old export column"
            workbook.save(path)
            workbook.close()
            row = ("=SYMBOL", 2, 600, 0.02, 12, 1200, 0.2, 200, 500, 1000)
            sync.write_positions([row], path)
            workbook = openpyxl.load_workbook(path)
            try:
                self.assertEqual(workbook["Other"]["A1"].value, "=1+1")
                for name in sync.SHEET_NAMES:
                    sheet = workbook[name]
                    self.assertEqual(tuple(c.value for c in sheet[1][:10]), sync.COLUMNS)
                    self.assertEqual(sheet["F2"].value, 1200)
                    self.assertEqual(sheet["A2"].data_type, "s")
                    self.assertEqual(sheet["D2"].number_format, "0.00%")
                    self.assertEqual(sheet["G2"].number_format, "0.00%")
                    self.assertEqual(sheet.auto_filter.ref, "A1:J2")
                    self.assertIsNone(sheet["A10"].value)
                    self.assertIsNone(sheet["N2"].value)
            finally:
                workbook.close()
            self.assertEqual(list(Path(directory).iterdir()), [path])


if __name__ == "__main__":
    unittest.main()
