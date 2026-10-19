"""Export positions and open orders to .xlsx, then open the saved workbook."""

from datetime import datetime, timedelta, timezone
import math
import os
from pathlib import Path
import tempfile

from dotenv import load_dotenv
import openpyxl
import requests
import schwabdev

from schwab_tokens import TOKENS_DB


MASTER_FILENAME = Path(__file__).resolve().parent / "data" / "Trading_Master_Dashboard.xlsx"
SHEET_NAMES = ("Raw_Positions", "Dashboard")
COLUMNS = (
    "Symbol/Description", "Qty", "Price", "Price Chng %", "Price Chng $",
    "Mkt Val", "Gain/Loss %", "Gain/Loss $", "Cost/Share", "Cost Basis",
)
SELL_SHEET = "Open_Sell_Orders"
BUY_SHEET = "Open_Buy_Orders"
SELL_COLUMNS = (
    "Account", "Symbol/Description", "Instruction", "Leg Qty", "Order Filled Qty",
    "Order Remaining Qty", "Order Type", "Limit Price", "Stop Price", "Status",
    "Duration", "Session", "Order ID", "Entered Time", "Strategy",
)


def fetch_open_orders(client):
    """Read 90 days of open BUY/SELL legs once, including nested orders."""
    response = client.linked_accounts()
    response.raise_for_status()
    accounts = response.json()
    if not isinstance(accounts, list) or not accounts:
        raise ValueError("Invalid linked accounts response; workbook not changed.")
    end = datetime.now(timezone.utc).replace(microsecond=0)
    beginning = end - timedelta(days=90)
    rows = {}
    terminal = {"FILLED", "CANCELED", "REJECTED", "EXPIRED", "REPLACED"}
    for account in accounts:
        if not isinstance(account, dict) or not account.get("hashValue"):
            raise ValueError("Linked account hash missing; workbook not changed.")
        window_end = end
        while window_end > beginning:
            window_start = max(beginning, window_end - timedelta(days=30))
            response = client.account_orders(
                accountHash=account["hashValue"],
                fromEnteredTime=window_start.isoformat(timespec="milliseconds"),
                toEnteredTime=window_end.isoformat(timespec="milliseconds"),
                maxResults=3000,
            )
            response.raise_for_status()
            orders = response.json()
            if not isinstance(orders, list) or len(orders) >= 3000:
                raise ValueError("Order response invalid or possibly truncated; workbook not changed.")
            pending = list(orders)
            while pending:
                order = pending.pop()
                if not isinstance(order, dict):
                    raise ValueError("Invalid order entry; workbook not changed.")
                pending.extend(order.get("childOrderStrategies", []))
                status = order.get("status")
                if status in terminal:
                    continue
                for index, leg in enumerate(order.get("orderLegCollection", [])):
                    instruction = leg.get("instruction", "")
                    if instruction not in {
                        "SELL", "SELL_SHORT", "SELL_TO_CLOSE", "SELL_TO_OPEN",
                        "BUY", "BUY_TO_COVER", "BUY_TO_CLOSE", "BUY_TO_OPEN",
                    }:
                        continue
                    instrument = leg.get("instrument", {})
                    symbol = instrument.get("symbol")
                    order_id = order.get("orderId")
                    if not symbol or order_id is None or not status:
                        raise ValueError("Order symbol, ID or status missing; workbook not changed.")
                    description = instrument.get("description")
                    label = f"{symbol} - {description}" if description else symbol
                    key = (account["hashValue"], str(order_id), index)
                    rows[key] = (
                        f"[...{str(account.get('accountNumber') or account['hashValue'])[-8:]}]",
                        label, instruction, number(leg.get("quantity")),
                        number(order.get("filledQuantity")),
                        number(order.get("remainingQuantity")),
                        order.get("orderType"), number(order.get("price")),
                        number(order.get("stopPrice")), status, order.get("duration"),
                        order.get("session"), str(order_id), order.get("enteredTime"),
                        order.get("orderStrategyType"),
                    )
            window_end = window_start
    return sorted(rows.values(), key=lambda row: (row[1], row[0], row[12]))


def number(value):
    if value is None:
        return None
    result = float(value)
    if not math.isfinite(result):
        raise ValueError("Schwab returned a non-finite numeric value; workbook not changed.")
    return result


def fetch_positions(client):
    """Validate the entire response before permitting a workbook update."""
    response = client.account_details_all(fields="positions")
    response.raise_for_status()
    accounts = response.json()
    if isinstance(accounts, dict):
        accounts = accounts.get("securitiesAccountList", [accounts])
    if not isinstance(accounts, list) or not accounts:
        raise ValueError("Schwab returned no valid account list; workbook not changed.")

    positions_to_export = []
    for item in accounts:
        if not isinstance(item, dict):
            raise ValueError("Invalid account entry; workbook not changed.")
        account = item.get("securitiesAccount", item)
        if not isinstance(account, dict):
            raise ValueError("Invalid securities account; workbook not changed.")
        identifier = account.get("accountNumber") or account.get("hashValue")
        if not identifier or "positions" not in account:
            raise ValueError("Account identifier or positions missing; workbook not changed.")
        positions = account["positions"]
        if not isinstance(positions, list):
            raise ValueError("Invalid positions list; workbook not changed.")

        for position in positions:
            if not isinstance(position, dict):
                raise ValueError("Invalid position entry; workbook not changed.")
            instrument = position.get("instrument")
            if not isinstance(instrument, dict) or not instrument.get("symbol"):
                raise ValueError("Position symbol missing; workbook not changed.")
            positions_to_export.append(position)

    symbols = sorted({p["instrument"]["symbol"] for p in positions_to_export})
    quotes = {}
    for start in range(0, len(symbols), 100):
        response = client.quotes(symbols=symbols[start:start + 100])
        response.raise_for_status()
        batch = response.json()
        if not isinstance(batch, dict):
            raise ValueError("Invalid quotes response; workbook not changed.")
        quotes.update(batch)

    rows = []
    for position in positions_to_export:
        instrument = position["instrument"]
        symbol = instrument["symbol"]
        long_qty = number(position.get("longQuantity", 0))
        short_qty = number(position.get("shortQuantity", 0))
        if long_qty is None or short_qty is None or (long_qty and short_qty):
            raise ValueError(f"Cannot calculate net position for {symbol}; workbook not changed.")
        quantity = long_qty - short_qty
        quote_entry = quotes.get(symbol, {})
        reference = quote_entry.get("reference", {}) if isinstance(quote_entry, dict) else {}
        description = instrument.get("description") or reference.get("description")
        label = f"{symbol} - {description}" if description else symbol
        quote = quote_entry.get("quote", {}) if isinstance(quote_entry, dict) else {}
        regular = quote_entry.get("regular", {}) if isinstance(quote_entry, dict) else {}
        if not isinstance(quote, dict) or not isinstance(regular, dict):
            raise ValueError(f"Invalid quote for {symbol}; workbook not changed.")
        if instrument.get("assetType") == "MUTUAL_FUND":
            price = number(quote.get("nAV", quote.get("lastPrice")))
            change_percent = number(quote.get("netPercentChange"))
            change_dollars = number(quote.get("netChange"))
        else:
            price = number(regular.get("regularMarketLastPrice"))
            change_percent = number(regular.get("regularMarketPercentChange"))
            change_dollars = number(regular.get("regularMarketNetChange"))
        if price is None or change_percent is None or change_dollars is None:
            print(f"Warning: Missing quote metrics for {symbol}; unavailable cells left blank.")
        market_value = average = None
        cost_basis = gain = gain_percent = None
        if instrument.get("assetType") in {"EQUITY", "COLLECTIVE_INVESTMENT", "MUTUAL_FUND"}:
            if price is not None:
                market_value = quantity * price
            gain_field = "shortOpenProfitLoss" if short_qty else "longOpenProfitLoss"
            position_gain = number(position.get(gain_field))
            position_value = number(position.get("marketValue"))
            # Preserve reported basis, but revalue P/L at the regular-session price.
            if position_gain is not None and position_value is not None:
                cost_basis = position_value - position_gain
                if quantity:
                    average = cost_basis / quantity
            else:
                basis_field = "taxLotAverageShortPrice" if short_qty else "taxLotAverageLongPrice"
                average = number(position.get(basis_field))
                if average is not None:
                    cost_basis = quantity * average
            if cost_basis is not None and market_value is not None:
                gain = market_value - cost_basis
                if cost_basis:
                    gain_percent = gain / abs(cost_basis)
            if cost_basis is None:
                print(f"Warning: Missing reported cost basis for {symbol}; cost basis and gain/loss left blank.")
        else:
            price = None
            print(f"Warning: Position valuation unsupported for {symbol}'s asset type; price, cost basis and gain/loss left blank.")
        rows.append((
            label, quantity, price,
            change_percent / 100 if change_percent is not None else None,
            change_dollars, market_value, gain_percent, gain, average, cost_basis,
        ))
    return sorted(
        rows,
        key=lambda row: (
            row[6] is None,
            -row[6] if row[6] is not None else 0,
            row[0],
        ),
    )


def write_positions(rows, filename=MASTER_FILENAME, *, sell_rows=None, buy_rows=None):
    """Replace target-sheet contents in a macro-free workbook, keeping other sheets."""
    filename = Path(filename)
    if filename.suffix.lower() != ".xlsx":
        raise ValueError("Position exports require a macro-free .xlsx workbook.")
    workbook = openpyxl.load_workbook(filename, keep_vba=False)
    temporary_path = None
    try:
        for name in SHEET_NAMES:
            sheet = workbook[name] if name in workbook.sheetnames else workbook.create_sheet(name)
            for cells in sheet:
                for cell in cells:
                    cell.value = None
            for row_number, values in enumerate([COLUMNS, *rows], start=1):
                for column, value in enumerate(values, start=1):
                    cell = sheet.cell(row_number, column, value)
                    if isinstance(value, str):
                        cell.data_type = "s"
                    if row_number > 1:
                        cell.number_format = (
                            '0.00%' if column in (4, 7)
                            else '$#,##0.00' if column in (3, 5, 6, 8, 9, 10)
                            else '#,##0.####' if column == 2 else "General"
                        )
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = f"A1:J{len(rows) + 1}"
            sheet.column_dimensions["A"].width = 55
            for column in "BCDEFGHIJ":
                sheet.column_dimensions[column].width = 18

        for name, order_rows in ((SELL_SHEET, sell_rows), (BUY_SHEET, buy_rows)):
            if order_rows is None:
                continue
            sheet = workbook[name] if name in workbook.sheetnames else workbook.create_sheet(name)
            for cells in sheet:
                for cell in cells:
                    cell.value = None
            for row_number, values in enumerate([SELL_COLUMNS, *order_rows], start=1):
                for column, value in enumerate(values, start=1):
                    cell = sheet.cell(row_number, column, value)
                    if isinstance(value, str):
                        cell.data_type = "s"
                    cell.number_format = (
                        '$#,##0.00' if row_number > 1 and column in (8, 9)
                        else '#,##0.####' if row_number > 1 and column in (4, 5, 6)
                        else "General"
                    )
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = f"A1:O{len(order_rows) + 1}"
            for column in "ABCDEFGHIJKLMNO":
                sheet.column_dimensions[column].width = 22
            sheet.column_dimensions["B"].width = 55
            sheet.column_dimensions["N"].width = 30

        # Save beside the original so a failed save never truncates the workbook.
        with tempfile.NamedTemporaryFile(
            dir=filename.parent, suffix=filename.suffix, delete=False
        ) as temporary:
            temporary_path = Path(temporary.name)
        workbook.save(temporary_path)
        workbook.close()
        os.replace(temporary_path, filename)
    finally:
        workbook.close()
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


def main():
    load_dotenv(Path(__file__).resolve().parent / ".env")
    credentials = [
        os.getenv("SCHWAB_CLIENT_ID"),
        os.getenv("SCHWAB_CLIENT_SECRET"),
        os.getenv("SCHWAB_CALLBACK_URL"),
    ]
    if not all(credentials):
        raise SystemExit("Error: Missing Schwab credentials in .env.")
    try:
        if not MASTER_FILENAME.is_file():
            raise FileNotFoundError(f"Workbook not found: {MASTER_FILENAME}")
        client = schwabdev.Client(*credentials, tokens_db=TOKENS_DB, timeout=60)
        print("Fetching current positions from Schwab...")
        rows = fetch_positions(client)
        print("Fetching open BUY/SELL orders entered within the last 90 days...")
        order_rows = fetch_open_orders(client)
        sell_rows = [row for row in order_rows if row[2].startswith("SELL")]
        buy_rows = [row for row in order_rows if row[2].startswith("BUY")]
        print(f"Retrieved {len(rows)} positions, {len(sell_rows)} open SELL legs and {len(buy_rows)} open BUY legs. Updating workbook...")
        write_positions(rows, sell_rows=sell_rows, buy_rows=buy_rows)
    except PermissionError as exc:
        raise SystemExit(
            "Error: Cannot access the workbook. Close it in Excel and retry; "
            "also check file permissions."
        ) from exc
    except (requests.RequestException, ValueError, OSError) as exc:
        raise SystemExit(f"Position sync failed: {exc}") from exc
    print(f"Saved {len(rows)} positions, {len(sell_rows)} open SELL legs and {len(buy_rows)} open BUY legs to {MASTER_FILENAME}")
    try:
        os.startfile(MASTER_FILENAME)
    except OSError as exc:
        raise SystemExit(
            f"Workbook saved, but could not open it automatically: {exc}"
        ) from exc


if __name__ == "__main__":
    main()