import os
import sys
from datetime import datetime, timedelta, timezone
import pandas as pd
from dotenv import load_dotenv
import schwabdev
import openpyxl
from openpyxl.utils.dataframe import dataframe_to_rows

# 1. Load system environment profiles
load_dotenv()
client_id = os.getenv("SCHWAB_CLIENT_ID")
client_secret = os.getenv("SCHWAB_CLIENT_SECRET")
callback_url = os.getenv("SCHWAB_CALLBACK_URL")

if not all([client_id, client_secret, callback_url]):
    raise ValueError("Error: Sensitive credential keys failed to load from your .env file.")

# 2. Initialize Schwab connection
try:
    client = schwabdev.Client(client_id, client_secret, callback_url)
except Exception as e:
    raise SystemExit(f"💥 Connection Failure: Schwab Client rejected authentication: {e}")


def sync_trading_engine_to_excel():
    master_filename = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "Trading_Master_Dashboard.xlsx")
    orders_sheet_name = "Raw_Orders"
    positions_sheet_name = "Raw_Positions"

    if not os.path.exists(master_filename):
        print(f"❌ Error: Could not find '{master_filename}'. Please create this file manually first!")
        return

    # A. FETCH DATA FROM SCHWAB (Pulling Account Details with Positions populated)
    try:
        print("🔍 Querying Schwab API for linked profiles and live balances...")
        accounts_raw = client.account_details_all(fields="positions")
        accounts_data = accounts_raw.json() if hasattr(accounts_raw, "json") else accounts_raw
    except Exception as e:
        print(f"⚠ API Error: Could not reach Schwab server: {e}")
        return

    if isinstance(accounts_data, list):
        accounts_list = accounts_data
    elif isinstance(accounts_data, dict):
        accounts_list = accounts_data.get("securitiesAccountList", [accounts_data])
    else:
        accounts_list = []

    print(f"ℹ Found {len(accounts_list)} account container(s) to scan.")
    
    all_orders_data = []
    all_positions_data = []
    
    # 60 days time window for working GTC orders
    now_utc = datetime.now(timezone.utc)
    start_utc = now_utc - timedelta(days=60)
    from_time_str = start_utc.isoformat(timespec='milliseconds').replace("+00:00", "Z")
    to_time_str = now_utc.isoformat(timespec='milliseconds').replace("+00:00", "Z")

    for account_item in accounts_list:
        if not isinstance(account_item, dict):
            continue

        sec_account = account_item.get("securitiesAccount", account_item)
        account_identifier = sec_account.get("hashValue") or sec_account.get("accountNumber")

        if not account_identifier:
            continue

        account_disp = f"[...{str(account_identifier)[-8:]}]"
        print(f"💼 Processing Account: {account_disp}")

        # ---------------------------------------------------------------------
        # STREAM 1: EXTRACT LIVE POSITIONS / INVENTORY HOLDINGS
        # ---------------------------------------------------------------------
        positions_list = sec_account.get("positions", [])
        print(f"   📊 Found {len(positions_list)} asset holding line items.")
        
        for pos in positions_list:
            if not isinstance(pos, dict):
                continue
                
            qty = float(pos.get("longQuantity", 0) or pos.get("shortQuantity", 0) or 0)
            avg_price = float(pos.get("averagePrice", 0.0) or pos.get("averageCost", 0.0) or 0.0)
            market_val = float(pos.get("marketValue", 0.0) or 0.0)
            
            current_price = float(pos.get("marketPrice", 0.0))
            if current_price == 0.0 and qty != 0:
                current_price = market_val / qty
            
            cost_basis = qty * avg_price
            gain_loss_dlr = market_val - cost_basis
            gain_loss_pct = (gain_loss_dlr / cost_basis) if cost_basis != 0 else 0.0
            
            instrument = pos.get("instrument", {})
            symbol = instrument.get("symbol", "N/A")
            
            all_positions_data.append({
                "Account": account_disp,
                "Symbol": symbol,
                "Quantity": qty,
                "Cost/Share": avg_price,
                "Average Price": avg_price,
                "Cost Basis": cost_basis,
                "Market Value": market_val,
                "Current Price": current_price,
                "Gain/Loss $": gain_loss_dlr,
                "Gain/Loss %": gain_loss_pct
            })

        # ---------------------------------------------------------------------
        # STREAM 2: EXTRACT ACTIVE WORKING ORDERS
        # ---------------------------------------------------------------------
        try:
            orders_raw = client.account_orders(
                accountHash=account_identifier, 
                fromEnteredTime=from_time_str, 
                toEnteredTime=to_time_str, 
                status="WORKING"
            )
            orders_list = orders_raw.json() if hasattr(orders_raw, "json") else orders_raw
        except Exception as order_err:
            print(f"   ⚠ Failed targeted orders sweep for {account_disp}: {order_err}")
            continue

        if isinstance(orders_list, dict):
            orders_list = orders_list.get("orderList") or orders_list.get("orders", [])
        if not isinstance(orders_list, list):
            orders_list = []

        for order in orders_list:
            if not isinstance(order, dict):
                continue

            order_id = order.get("orderId", "N/A")
            entered_time = order.get("enteredTime", "")
            order_type = order.get("orderType", "")
            price = order.get("price", 0) or order.get("stopPrice", 0)

            for leg in order.get("orderLegCollection", []):
                instruction = leg.get("instruction", "")
                
                if "SELL" in instruction.upper():
                    leg_instrument = leg.get("instrument", {})
                    symbol = leg_instrument.get("symbol", "N/A")
                    qty = leg.get("quantity", 0)

                    all_orders_data.append({
                        "Account": account_disp,
                        "Order_ID": str(order_id),
                        "Symbol": symbol,
                        "Instruction": instruction,
                        "Quantity": qty,
                        "Order_Type": order_type,
                        "Execution_Price": float(price) if price else 0.0,
                        "Time_Stamp": entered_time[:10] if entered_time else "N/A"
                    })

    # Convert to standard DataFrames for unified sorting
    df_orders = pd.DataFrame(all_orders_data)
    if not df_orders.empty:
        df_orders = df_orders.sort_values(by=["Account", "Symbol"])
    else:
        df_orders = pd.DataFrame(columns=["Account", "Order_ID", "Symbol", "Instruction", "Quantity", "Order_Type", "Execution_Price", "Time_Stamp"])

    df_positions = pd.DataFrame(all_positions_data)
    if not df_positions.empty:
        df_positions = df_positions.sort_values(by=["Account", "Symbol"])
    else:
        df_positions = pd.DataFrame(columns=["Account", "Symbol", "Quantity", "Cost/Share", "Average Price", "Cost Basis", "Market Value", "Current Price", "Gain/Loss $", "Gain/Loss %"])

    # B. DUAL-STREAM IN-PLACE EXCEL INJECTION
    try:
        print(f"\n💾 Opening '{master_filename}' to update sheets...")
        workbook = openpyxl.load_workbook(master_filename)
        
        # --- Update Raw Orders Tab ---
        if orders_sheet_name not in workbook.sheetnames:
            workbook.create_sheet(title=orders_sheet_name)
        ws_orders = workbook[orders_sheet_name]
        ws_orders.delete_rows(1, ws_orders.max_row + 100)
        for r in dataframe_to_rows(df_orders, index=False, header=True):
            ws_orders.append(r)
        for row_num in range(2, ws_orders.max_row + 1):
            ws_orders.cell(row=row_num, column=5).number_format = '#,##0'  # Quantity
            ws_orders.cell(row=row_num, column=7).number_format = '$#,##0.00'  # Price

        # --- Update Raw Positions Tab ---
        if positions_sheet_name not in workbook.sheetnames:
            workbook.create_sheet(title=positions_sheet_name)
        ws_positions = workbook[positions_sheet_name]
        ws_positions.delete_rows(1, ws_positions.max_row + 100)
        for r in dataframe_to_rows(df_positions, index=False, header=True):
            ws_positions.append(r)
            
        for row_num in range(2, ws_positions.max_row + 1):
            ws_positions.cell(row=row_num, column=3).number_format = '#,##0.00'   # Quantity
            ws_positions.cell(row=row_num, column=4).number_format = '$#,##0.00'  # Cost/Share
            ws_positions.cell(row=row_num, column=5).number_format = '$#,##0.00'  # Average Price
            ws_positions.cell(row=row_num, column=6).number_format = '$#,##0.00'  # Cost Basis
            ws_positions.cell(row=row_num, column=7).number_format = '$#,##0.00'  # Market Value
            ws_positions.cell(row=row_num, column=8).number_format = '$#,##0.00'  # Current Price
            ws_positions.cell(row=row_num, column=9).number_format = '$#,##0.00'  # Gain/Loss $
            ws_positions.cell(row=row_num, column=10).number_format = '0.00%'     # Gain/Loss %

        workbook.save(master_filename)
        print(f"🔄 Success: '{orders_sheet_name}' and '{positions_sheet_name}' are synchronized perfectly at {datetime.now().strftime('%H:%M:%S')}!")
        
    except PermissionError:
        print(f"❌ Excel Sync Failed: The file '{master_filename}' is open in Excel. Close it and run the script again!")
    except Exception as ex:
        print(f"❌ Workbook Manipulation Error: {ex}")


if __name__ == "__main__":
    print("🚀 Launching Self-Hosted Double-Stream Excel Dashboard Engine...")
    sync_trading_engine_to_excel()
