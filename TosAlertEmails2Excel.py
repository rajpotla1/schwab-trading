"""
Read today's thinkorswim alert emails and export them to an Excel file.
Only subjects containing the word BOT followed by @ are exported. BOT Text
contains the text from BOT up to, but excluding, the first following @.
Matching is case-insensitive; the full Subject is preserved.
The Trades worksheet lists Date (received), Ticker, Qty, Price (number
after @) and New (Yes = email was not in the previous export of the same
output file; new rows are highlighted) for each exported email,
in the same newest-first received-time order as Alerts. Unrecognized BOT trade formats
stop the export rather than silently omitting or guessing a quantity.
Classic Outlook searches all mail folders in the selected mailbox, including
top-level folders outside the Inbox. Folder access errors stop the export.
After the Excel file is saved, the exported BOT emails are marked as read
(classic Outlook only; skip with --no-mark-read).

Sources:
  graph   - Hotmail/Outlook.com directly via Microsoft Graph (works with "New Outlook"
            or no Outlook at all). Requires OUTLOOK_CLIENT_ID in .env (see README note below).
  outlook - Classic Outlook desktop app via COM (New Outlook is NOT supported).

Graph setup (one time):
  1. https://entra.microsoft.com -> App registrations -> New registration
     Supported account types: "Personal Microsoft accounts only"
  2. Authentication -> Add platform "Mobile and desktop applications",
     tick https://login.microsoftonline.com/common/oauth2/nativeclient,
     and set "Allow public client flows" = Yes.
  3. Copy the Application (client) ID into .env as OUTLOOK_CLIENT_ID=...
  First run prints a code to enter at https://microsoft.com/devicelogin; the token is
  cached in Config/outlook_token.json afterwards.

Usage (set OUTLOOK_ACCOUNT=you@hotmail.com in .env to pick the mailbox):
    python TosAlertEmails2Excel.py                      # BOT ... @ alerts today (classic Outlook)
    python TosAlertEmails2Excel.py --subject "Alert"    # additional subject filter
    python TosAlertEmails2Excel.py --source graph       # Microsoft Graph instead of Outlook
"""
import argparse
import os
import re
from datetime import datetime, timedelta, timezone

import pandas as pd
import requests
from dotenv import load_dotenv
from openpyxl.styles import PatternFill

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
SENDER = "alerts@thinkorswim.com"
GRAPH_SCOPES = ["Mail.Read"]
TOKEN_CACHE = os.path.join(SCRIPT_DIR, "Config", "outlook_token.json")
OL_MAIL_ITEM_CLASS = 43


def extract_bot_text(subject: str) -> str | None:
    match = re.search(r"\bBOT\b[^@]*@", subject, flags=re.IGNORECASE)
    return match.group(0)[:-1].strip() if match else None


def parse_bot_trade(bot_text: str) -> tuple[str, int | float]:
    match = re.fullmatch(
        r"BOT\s+\+?(\d+(?:\.\d+)?)\s+(\S+)", bot_text, flags=re.IGNORECASE
    )
    if match is None:
        raise ValueError(f"Cannot extract Qty and Ticker from BOT text: {bot_text!r}")
    quantity, ticker = match.groups()
    return ticker.upper(), float(quantity) if "." in quantity else int(quantity)


def extract_bot_price(subject: str) -> float:
    match = re.search(r"\bBOT\b[^@]*@\s*(\d+(?:\.\d+)?)", subject, flags=re.IGNORECASE)
    if match is None:
        raise ValueError(f"Cannot extract Price from subject: {subject!r}")
    return float(match.group(1))


def get_sender_address(item):
    try:
        if item.SenderEmailType == "EX":
            exch_user = item.Sender.GetExchangeUser()
            if exch_user:
                return exch_user.PrimarySmtpAddress
        return item.SenderEmailAddress or ""
    except Exception:
        return ""


def iter_folders(folder):
    if folder.DefaultItemType == 0:
        yield folder
    for sub in folder.Folders:
        yield from iter_folders(sub)


def get_mailbox_roots(namespace, account):
    roots = []
    for store in namespace.Stores:
        if account and account.lower() not in store.DisplayName.lower():
            continue
        if store.ExchangeStoreType == 2:  # Exchange public folders are not mailboxes.
            continue
        roots.append(store.GetRootFolder())
    return roots


def get_graph_token(client_id):
    import msal

    cache = msal.SerializableTokenCache()
    if os.path.exists(TOKEN_CACHE):
        with open(TOKEN_CACHE, "r") as f:
            cache.deserialize(f.read())
    app = msal.PublicClientApplication(
        client_id, authority="https://login.microsoftonline.com/consumers", token_cache=cache
    )
    result = None
    accounts = app.get_accounts()
    if accounts:
        result = app.acquire_token_silent(GRAPH_SCOPES, account=accounts[0])
    if not result:
        flow = app.initiate_device_flow(scopes=GRAPH_SCOPES)
        if "user_code" not in flow:
            raise SystemExit(f"Could not start sign-in: {flow.get('error_description', flow)}")
        print(flow["message"], flush=True)
        result = app.acquire_token_by_device_flow(flow)
    if "access_token" not in result:
        raise SystemExit(f"Sign-in failed: {result.get('error_description', result)}")
    if cache.has_state_changed:
        os.makedirs(os.path.dirname(TOKEN_CACHE), exist_ok=True)
        with open(TOKEN_CACHE, "w") as f:
            f.write(cache.serialize())
    return result["access_token"]


def collect_alerts_graph(day, subject_filter):
    load_dotenv(os.path.join(SCRIPT_DIR, ".env"))
    client_id = os.getenv("OUTLOOK_CLIENT_ID")
    if not client_id:
        raise SystemExit("OUTLOOK_CLIENT_ID is missing from .env. See setup steps at the top of this script.")
    token = get_graph_token(client_id)

    local_start = datetime(day.year, day.month, day.day).astimezone()
    start_utc = local_start.astimezone(timezone.utc)
    end_utc = (local_start + timedelta(days=1)).astimezone(timezone.utc)
    url = "https://graph.microsoft.com/v1.0/me/messages"
    params = {
        "$filter": (
            f"receivedDateTime ge {start_utc:%Y-%m-%dT%H:%M:%SZ} and "
            f"receivedDateTime lt {end_utc:%Y-%m-%dT%H:%M:%SZ}"
        ),
        "$select": "receivedDateTime,from,subject,body,parentFolderId",
        "$top": "100",
    }
    headers = {"Authorization": f"Bearer {token}", "Prefer": 'outlook.body-content-type="text"'}

    rows = []
    while url:
        resp = requests.get(url, headers=headers, params=params, timeout=60)
        if not resp.ok:
            raise SystemExit(f"Graph request failed ({resp.status_code}): {resp.text}")
        data = resp.json()
        for msg in data.get("value", []):
            sender = (msg.get("from") or {}).get("emailAddress", {}).get("address", "")
            if sender.lower() != SENDER:
                continue
            subject = msg.get("subject") or ""
            if subject_filter and subject_filter.lower() not in subject.lower():
                continue
            received = datetime.fromisoformat(msg["receivedDateTime"].replace("Z", "+00:00"))
            rows.append({
                "Received": received.astimezone().replace(tzinfo=None),
                "From": sender,
                "Subject": subject,
                "Body": re.sub(r"\s+", " ", (msg.get("body") or {}).get("content", "")).strip(),
                "Folder": msg.get("parentFolderId", ""),
            })
        url = data.get("@odata.nextLink")
        params = None
    return rows


def collect_alerts(day, subject_filter, account):
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    outlook = win32com.client.Dispatch("Outlook.Application")
    namespace = outlook.GetNamespace("MAPI")

    start = datetime(day.year, day.month, day.day)
    end = start + timedelta(days=1)
    date_filter = (
        f"[ReceivedTime] >= '{start.strftime('%m/%d/%Y %I:%M %p')}' AND "
        f"[ReceivedTime] < '{end.strftime('%m/%d/%Y %I:%M %p')}'"
    )

    roots = get_mailbox_roots(namespace, account)
    if not roots:
        raise SystemExit(f"No Outlook mailbox found{' for account ' + account if account else ''}.")

    rows = []
    for root in roots:
        for folder in iter_folders(root):
            items = folder.Items.Restrict(date_filter)
            for item in items:
                if getattr(item, "Class", None) != OL_MAIL_ITEM_CLASS:
                    continue
                if get_sender_address(item).lower() != SENDER:
                    continue
                subject = item.Subject or ""
                if subject_filter and subject_filter.lower() not in subject.lower():
                    continue
                received = item.ReceivedTime
                rows.append({
                    "Received": datetime(received.year, received.month, received.day,
                                         received.hour, received.minute, received.second),
                    "From": SENDER,
                    "Subject": subject,
                    "Body": " ".join((item.Body or "").split()),
                    "Folder": folder.FolderPath,
                    "EntryID": item.EntryID,
                    "StoreID": folder.StoreID,
                })
    return rows


def mark_emails_read(rows):
    import win32com.client

    namespace = win32com.client.Dispatch("Outlook.Application").GetNamespace("MAPI")
    marked = 0
    for row in rows:
        item = namespace.GetItemFromID(row["EntryID"], row["StoreID"])
        if item.UnRead:
            item.UnRead = False
            item.Save()
            marked += 1
    return marked


def email_key(received, subject):
    return pd.Timestamp(received).floor("s"), subject


def load_previous_keys(path):
    if not os.path.exists(path) or os.path.getsize(path) == 0:
        return set()
    previous = pd.read_excel(path, sheet_name="Alerts", usecols=["Received", "Subject"])
    return {email_key(r, s) for r, s in zip(previous["Received"], previous["Subject"].fillna(""))}


def main():
    load_dotenv(os.path.join(SCRIPT_DIR, ".env"))
    parser = argparse.ArgumentParser(description="Export thinkorswim alert emails to Excel.")
    parser.add_argument("--source", choices=["graph", "outlook"], default="outlook",
                        help="outlook = classic Outlook desktop (default); graph = Hotmail via Microsoft Graph.")
    parser.add_argument("--subject", default="", help="Additional subject text filter; BOT followed by @ is always required.")
    parser.add_argument("--account", default=os.getenv("OUTLOOK_ACCOUNT", ""),
                        help="Outlook account/store name to search (default: OUTLOOK_ACCOUNT in .env; "
                             "all mailboxes if unset).")
    parser.add_argument("--date", default="", help="Date to pull (YYYY-MM-DD). Defaults to today.")
    parser.add_argument("--out", default="", help="Output Excel path (default: data\\TdyBuyAlertsTrades.xlsx).")
    parser.add_argument("--no-open", action="store_true", help="Do not open the Excel file after saving.")
    parser.add_argument("--no-mark-read", action="store_true",
                        help="Do not mark the exported BOT emails as read (classic Outlook source only).")
    args = parser.parse_args()

    day = datetime.strptime(args.date, "%Y-%m-%d") if args.date else datetime.now()
    if args.source == "graph":
        rows = collect_alerts_graph(day, args.subject)
    else:
        rows = collect_alerts(day, args.subject, args.account)

    matching_rows = []
    for row in rows:
        bot_text = extract_bot_text(row["Subject"])
        if bot_text is not None:
            matching_rows.append({**row, "BOT Text": bot_text})

    out_path = os.path.abspath(args.out or os.path.join(SCRIPT_DIR, "data", "TdyBuyAlertsTrades.xlsx"))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    try:
        with open(out_path, "ab"):
            pass
    except PermissionError:
        raise SystemExit(f"Cannot write {out_path}. Close it in Excel and run again.")
    previous_keys = load_previous_keys(out_path)

    df = pd.DataFrame(matching_rows, columns=["Received", "From", "Subject", "BOT Text", "Body", "Folder"])
    df = df.sort_values("Received", ascending=False, kind="stable")
    is_new = [email_key(r, s) not in previous_keys for r, s in zip(df["Received"], df["Subject"])]
    trades = pd.DataFrame(
        [
            (received, *parse_bot_trade(text), extract_bot_price(subject), "Yes" if new else "No")
            for received, text, subject, new in zip(df["Received"], df["BOT Text"], df["Subject"], is_new)
        ],
        columns=["Date", "Ticker", "Qty", "Price", "New"],
    )
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        trades.to_excel(writer, index=False, sheet_name="Trades")
        trades_ws = writer.sheets["Trades"]
        trades_ws.auto_filter.ref = trades_ws.dimensions
        trades_ws.freeze_panes = "A2"
        for col, width in zip("ABCDE", (20, 15, 10, 12, 8)):
            trades_ws.column_dimensions[col].width = width
        for cell in trades_ws["A"][1:]:
            cell.number_format = "yyyy-mm-dd hh:mm:ss"
        new_fill = PatternFill(fill_type="solid", start_color="FFFF99", end_color="FFFF99")
        for row_cells, new in zip(trades_ws.iter_rows(min_row=2), is_new):
            if new:
                for cell in row_cells:
                    cell.fill = new_fill
        df.to_excel(writer, index=False, sheet_name="Alerts")
        ws = writer.sheets["Alerts"]
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
        for col, width in zip("ABCDEF", (20, 26, 60, 60, 100, 40)):
            ws.column_dimensions[col].width = width
        for cell in ws["A"][1:]:
            cell.number_format = "yyyy-mm-dd hh:mm:ss"
        writer.book.active = 0

    print(f"Found {len(df)} email(s) from {SENDER} on {day:%Y-%m-%d}"
          " with BOT followed by @ in the subject"
          f"{f' with subject containing {args.subject!r}' if args.subject else ''}.")
    print(f"Saved to {out_path}")
    if not args.no_mark_read:
        if args.source == "outlook":
            print(f"Marked {mark_emails_read(matching_rows)} unread exported email(s) as read.")
        else:
            print("Mark-as-read is only supported with --source outlook; emails left unchanged.")
    if not args.no_open:
        os.startfile(out_path)


if __name__ == "__main__":
    main()
