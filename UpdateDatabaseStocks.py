"""
Apply new BOT trades from TdyBuyAlertsTrades.xlsx to the IRA sheet of
Database Stocks.xlsm (open in Excel), replicating the workbook's VBA macros:

  Ticker found     -> Button_today_purch_dt: B = trade date, clear highlighting
                      A:BF, T formula += Qty*Price, V formula += Qty.
  Ticker not found -> Sheet1.InsertRowFromSelection with row 3 selected: copy
                      row 3 to a new row 4, B/C/D = trade date, AM = ticker
                      (Stocks data type), T = =Qty*Price, V = =Qty.
Every added or updated row is then highlighted red across A:BF.

Only Trades rows with New = Yes that are not yet Applied are used; each applied
row is marked Applied = Yes in TdyBuyAlertsTrades.xlsx so reruns never
double-count. Tickers match Column A ("TICKER - TDIRA ()"); when a ticker has
several rows, the one with the latest Column B date is updated. The workbook
is NOT saved; review the changes and save it yourself.

Usage:
    python UpdateDatabaseStocks.py --dry-run   # show what would change
    python UpdateDatabaseStocks.py
"""
import argparse
import os
import re
import time
from datetime import datetime

import pandas as pd
from dotenv import load_dotenv
from openpyxl import load_workbook

from TosAlertEmails2Excel import SCRIPT_DIR, close_workbook_if_open, is_same_workbook

SHEET = "IRA"
TEMPLATE_ROW = 3
FIRST_TICKER_ROW = 3
TICKER_IN_A = re.compile(r"^([A-Z][A-Z0-9./ ]*?) - ")
XL_SHIFT_DOWN = -4121
XL_COLOR_INDEX_NONE = -4142
RED = 255  # Excel colors are BGR: RGB(255, 0, 0)
STOCKS_SERVICE_ID = 268435456
RPC_E_CALL_REJECTED = -2147418111
RPC_E_SERVERCALL_RETRYLATER = -2147417846


def normalize_ticker(ticker):
    return re.sub(r"[./]", " ", str(ticker).strip().upper())


def fmt_num(value):
    text = f"{float(value):.6f}".rstrip("0").rstrip(".")
    return text or "0"


def append_term(formula, term):
    formula = str(formula or "")
    return f"={term}" if formula in ("", "=") else f"{formula}+{term}"


def load_new_trades(path):
    trades = pd.read_excel(path, sheet_name="Trades")
    if "Applied" not in trades.columns:
        trades["Applied"] = ""
    pending = trades[(trades["New"] == "Yes") & (trades["Applied"].fillna("") != "Yes")]
    return pending.sort_values("Date", kind="stable")


def com(action, attempts=60, delay=0.5):
    """Run one Excel COM step, retrying while Excel is busy (call rejected => not executed)."""
    import pywintypes

    for attempt in range(attempts):
        try:
            return action()
        except pywintypes.com_error as exc:
            if exc.hresult not in (RPC_E_CALL_REJECTED, RPC_E_SERVERCALL_RETRYLATER) or attempt == attempts - 1:
                raise
            time.sleep(delay)


def build_ticker_rows(ws):
    used = com(lambda: ws.UsedRange)
    last_row = com(lambda: used.Row + used.Rows.Count - 1)
    a_values = com(lambda: ws.Range(f"A{FIRST_TICKER_ROW}:A{last_row}").Value)
    b_values = com(lambda: ws.Range(f"B{FIRST_TICKER_ROW}:B{last_row}").Value)
    best = {}
    for offset, ((a,), (b,)) in enumerate(zip(a_values, b_values)):
        match = TICKER_IN_A.match(a) if isinstance(a, str) else None
        if not match:
            continue
        row = FIRST_TICKER_ROW + offset
        key = normalize_ticker(match.group(1))
        b_date = b if hasattr(b, "year") else None
        current = best.get(key)
        if current is None or (b_date is not None and (current[1] is None or b_date > current[1])):
            best[key] = (row, b_date)
    return {key: row for key, (row, _) in best.items()}


def set_value(ws, address, value):
    com(lambda: setattr(ws.Range(address), "Value", value))


def set_formula(ws, address, formula):
    com(lambda: setattr(ws.Range(address), "Formula", formula))


def get_formula(ws, address):
    return com(lambda: ws.Range(address).Formula)


def update_existing(ws, row, trade_date, qty, price):
    term = f"{fmt_num(qty)}*{fmt_num(price)}"
    set_value(ws, f"B{row}", trade_date)
    com(lambda: setattr(ws.Range(f"A{row}:BF{row}").Interior, "ColorIndex", XL_COLOR_INDEX_NONE))
    set_formula(ws, f"T{row}", append_term(get_formula(ws, f"T{row}"), term))
    set_formula(ws, f"V{row}", append_term(get_formula(ws, f"V{row}"), fmt_num(qty)))


def insert_new(ws, ticker, trade_date, qty, price):
    new_row = TEMPLATE_ROW + 1
    com(lambda: ws.Rows(new_row).Insert(Shift=XL_SHIFT_DOWN))
    com(lambda: ws.Rows(TEMPLATE_ROW).Copy(ws.Rows(new_row)))
    for col in "BCD":
        set_value(ws, f"{col}{new_row}", trade_date)
    set_formula(ws, f"V{new_row}", f"={fmt_num(qty)}")
    set_formula(ws, f"T{new_row}", f"={fmt_num(qty)}*{fmt_num(price)}")
    com(lambda: ws.Range(f"AM{new_row}").ClearContents())
    set_value(ws, f"AM{new_row}", ticker)
    try:
        com(lambda: ws.Range(f"AM{new_row}").ConvertToLinkedDataType(
            ServiceID=STOCKS_SERVICE_ID, LanguageCulture="en-US"))
    except Exception as exc:
        print(f"  Warning: could not convert {ticker} in AM{new_row} to a Stocks data type: {exc}")
    return new_row


def highlight_row(ws, row):
    com(lambda: setattr(ws.Range(f"A{row}:BF{row}").Interior, "Color", RED))


def apply_trades(ws, trades, dry_run=False, applied=None):
    applied = [] if applied is None else applied
    rows = build_ticker_rows(ws)
    for index, trade in trades.iterrows():
        ticker = str(trade["Ticker"]).upper()
        key = normalize_ticker(ticker)
        trade_date = datetime(trade["Date"].year, trade["Date"].month, trade["Date"].day)
        qty, price = trade["Qty"], trade["Price"]
        if key in rows:
            row = rows[key]
            print(f"UPDATE {ticker:<6} row {row}: B={trade_date:%m/%d/%Y}, T += {fmt_num(qty)}*{fmt_num(price)}, V += {fmt_num(qty)}")
            if not dry_run:
                update_existing(ws, row, trade_date, qty, price)
                highlight_row(ws, row)
        else:
            print(f"INSERT {ticker:<6} new row {TEMPLATE_ROW + 1} (copy of row {TEMPLATE_ROW}): "
                  f"T = {fmt_num(qty)}*{fmt_num(price)}, V = {fmt_num(qty)}")
            if not dry_run:
                new_row = insert_new(ws, ticker, trade_date, qty, price)
                highlight_row(ws, new_row)
                rows = {k: (r + 1 if r >= new_row else r) for k, r in rows.items()}
                rows[key] = new_row
            else:
                rows[key] = TEMPLATE_ROW + 1
        applied.append(index)
    return applied


def mark_applied(path, indexes):
    if not indexes:
        return
    wb = load_workbook(path)
    ws = wb["Trades"]
    headers = [cell.value for cell in ws[1]]
    if "Applied" in headers:
        col = headers.index("Applied") + 1
    else:
        col = len(headers) + 1
        ws.cell(row=1, column=col, value="Applied")
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 10
        ws.auto_filter.ref = ws.dimensions
    for index in indexes:
        ws.cell(row=index + 2, column=col, value="Yes")
    wb.save(path)


def get_workbook(path):
    import win32com.client

    try:
        excel = win32com.client.GetActiveObject("Excel.Application")
    except Exception:
        excel = win32com.client.Dispatch("Excel.Application")
    for workbook in excel.Workbooks:
        if is_same_workbook(workbook.FullName, path):
            return workbook
    if not os.path.exists(path):
        raise SystemExit(f"Workbook not found: {path}")
    excel.Visible = True
    return excel.Workbooks.Open(path)


def main():
    load_dotenv(os.path.join(SCRIPT_DIR, ".env"))
    default_workbook = os.getenv("DATABASE_STOCKS_PATH") or os.path.join(
        os.environ.get("OneDriveConsumer") or os.path.expanduser(r"~\OneDrive"), "Documents", "Database Stocks.xlsm")
    parser = argparse.ArgumentParser(description="Apply new BOT trades to Database Stocks.xlsm (IRA sheet).")
    parser.add_argument("--trades", default=os.path.join(SCRIPT_DIR, "data", "TdyBuyAlertsTrades.xlsx"))
    parser.add_argument("--workbook", default=default_workbook,
                        help="Database Stocks workbook (default: DATABASE_STOCKS_PATH in .env or OneDrive\\Documents).")
    parser.add_argument("--dry-run", action="store_true", help="Show the planned changes without modifying anything.")
    args = parser.parse_args()

    trades_path = os.path.abspath(args.trades)
    reopen_trades = close_workbook_if_open(trades_path)
    trades = load_new_trades(trades_path)
    if trades.empty:
        print("No new (unapplied) trades to apply.")
    else:
        ws = get_workbook(os.path.abspath(args.workbook)).Worksheets(SHEET)
        applied = []
        try:
            apply_trades(ws, trades, dry_run=args.dry_run, applied=applied)
        finally:
            if not args.dry_run:
                mark_applied(trades_path, applied)
        if args.dry_run:
            print(f"Dry run: {len(applied)} trade(s) would be applied. Nothing was changed.")
        else:
            print(f"Applied {len(applied)} trade(s) to {SHEET}. Database Stocks.xlsm is NOT saved; review and save it.")
    if reopen_trades:
        os.startfile(trades_path)


if __name__ == "__main__":
    main()
